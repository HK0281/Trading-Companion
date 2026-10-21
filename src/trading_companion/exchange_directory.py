"""Explicit, bounded reads of current SSE and SZSE security-name lists."""

from datetime import datetime, timezone
from io import BytesIO
import json
from math import isfinite
import re
import warnings
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from trading_companion.errors import DataSourceError, DataValidationError, InvalidRequestError
from trading_companion.security_directory import Security, SecurityDirectory


SSE_URL = "https://query.sse.com.cn/sseQuery/commonQuery.do"
SZSE_URL = "https://www.szse.cn/api/report/ShowReport"
SSE_QUERY = "COMMON_SSE_CP_GPJCTPZ_GPLB_GP_L"
DIRECTORY_SOURCE = "sse_main_star_including_cdr_and_szse_a_export_v1"
MAX_HTTP_BYTES = 10_000_000
MAX_XLSX_BYTES = 32_000_000
MAX_SECURITIES = 10_000


def _timeout(value: float) -> None:
    if (isinstance(value, bool) or not isinstance(value, (int, float))
            or not isfinite(value) or value <= 0):
        raise InvalidRequestError("Directory timeout must be a finite positive number")


def _download(url: str, params: dict[str, str], *, referer: str, timeout: float) -> bytes:
    """Do not redirect, retry, change providers, or return a partial body."""
    try:
        import requests
    except ImportError:
        raise DataSourceError("Install the data extra to refresh the security directory") from None
    try:
        response = requests.get(url, params=params,
                                headers={"User-Agent": "Mozilla/5.0", "Referer": referer},
                                timeout=(timeout, timeout), stream=True, allow_redirects=False)
        try:
            response.raise_for_status()
            if response.status_code != 200:
                raise DataSourceError("Directory endpoint did not return HTTP 200")
            chunks = []
            size = 0
            for chunk in response.iter_content(chunk_size=65_536):
                size += len(chunk)
                if size > MAX_HTTP_BYTES:
                    raise DataValidationError("Directory response is too large")
                chunks.append(chunk)
            if not size:
                raise DataValidationError("Directory response is empty")
            return b"".join(chunks)
        finally:
            response.close()
    except requests.exceptions.RequestException:
        raise DataSourceError("Exchange directory request failed; no snapshot was refreshed") from None


def _security(code: object, name: object, *, exchange: str, prefixes: tuple[str, ...]) -> Security:
    # Never turn missing, numeric, fractional, or malformed codes into valid-looking identities.
    if type(code) is not str or re.fullmatch(r"[0-9]{6}", code) is None or not code.startswith(prefixes):
        raise DataValidationError("Exchange directory returned an invalid security code")
    return Security(f"{code}.{exchange}", name)


def _checked_records(records: list[Security]) -> tuple[Security, ...]:
    if not 1 <= len(records) <= MAX_SECURITIES:
        raise DataValidationError("Exchange directory has an invalid record count")
    if len({record.instrument_id for record in records}) != len(records):
        raise DataValidationError("Exchange directory contains duplicate instruments")
    return tuple(records)


def parse_sse_directory(body: bytes, *, stock_type: str) -> tuple[Security, ...]:
    """Require the complete declared single page, not merely some valid rows."""
    if stock_type not in ("1", "8"):
        raise InvalidRequestError("Choose SSE main-board A shares or STAR shares")
    try:
        payload = json.loads(body.decode("utf-8"))
    except (UnicodeError, ValueError, RecursionError):
        raise DataValidationError("SSE directory returned invalid JSON") from None
    if type(payload) is not dict or payload.get("sqlId") != SSE_QUERY:
        raise DataValidationError("Unexpected SSE directory schema")
    if payload.get("actionErrors") != [] or payload.get("fieldErrors") != {}:
        raise DataValidationError("SSE directory reported errors or changed its error schema")
    page = payload.get("pageHelp")
    rows = payload.get("result")
    if (type(page) is not dict or type(rows) is not list
            or type(page.get("total")) is not int
            or not 1 <= page["total"] <= MAX_SECURITIES
            or type(page.get("pageNo")) is not int or page["pageNo"] != 1
            or type(page.get("pageCount")) is not int or page["pageCount"] != 1
            or len(rows) != page["total"]):
        raise DataValidationError("SSE directory is incomplete or needs pagination support")
    records = []
    for row in rows:
        if type(row) is not dict or row.get("STOCK_TYPE") != stock_type:
            raise DataValidationError("Unexpected SSE directory stock type")
        records.append(_security(row.get("A_STOCK_CODE"), row.get("SEC_NAME_CN"), exchange="SH",
                                 prefixes=("688", "689") if stock_type == "8" else ("600", "601", "603", "605")))
    return _checked_records(records)


def parse_szse_directory(body: bytes) -> tuple[Security, ...]:
    """Validate the exported table; it provides no independent total-count proof."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        raise DataSourceError("Install the data extra to read the SZSE directory export") from None
    try:
        with ZipFile(BytesIO(body)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(entry.file_size for entry in entries) > MAX_XLSX_BYTES:
                raise DataValidationError("SZSE workbook exceeds expansion limits")
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Workbook contains no default style.*", category=UserWarning)
            workbook = load_workbook(BytesIO(body), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) != 1:
                raise DataValidationError("Unexpected SZSE workbook sheets")
            sheet = workbook.worksheets[0]
            # The real export declares A1:A1 even when thousands of rows exist.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(max_col=100, values_only=True)
            header = next(rows, ())
            if any(header.count(field) != 1 for field in ("A股代码", "A股简称")):
                raise DataValidationError("SZSE directory is missing unique code/name columns")
            code_index, name_index = header.index("A股代码"), header.index("A股简称")
            records = []
            for index, row in enumerate(rows, start=1):
                if index > MAX_SECURITIES:
                    raise DataValidationError("SZSE directory has too many rows")
                records.append(_security(row[code_index], row[name_index], exchange="SZ",
                                         prefixes=("0", "3")))
            return _checked_records(records)
        finally:
            workbook.close()
    except (BadZipFile, InvalidFileException, ParseError, OSError, ValueError, TypeError, KeyError) as error:
        if isinstance(error, DataValidationError):
            raise
        raise DataValidationError("SZSE returned an invalid directory workbook") from None


class ExchangeDirectorySource:
    """Three required fragments, with no partial success or inferred aliases."""

    def __init__(self, *, timeout_seconds: float = 15.0):
        _timeout(timeout_seconds)
        self.timeout_seconds = timeout_seconds

    def read_directory(self) -> SecurityDirectory:
        records = []
        for stock_type in ("1", "8"):
            params = {
                "STOCK_TYPE": stock_type, "REG_PROVINCE": "", "CSRC_CODE": "", "STOCK_CODE": "",
                "sqlId": SSE_QUERY, "COMPANY_STATUS": "2,4,5,7,8", "type": "inParams",
                "isPagination": "true", "pageHelp.cacheSize": "1", "pageHelp.beginPage": "1",
                "pageHelp.pageSize": "10000", "pageHelp.pageNo": "1", "pageHelp.endPage": "1",
            }
            body = _download(SSE_URL, params, referer="https://www.sse.com.cn/assortment/stock/list/share/",
                             timeout=self.timeout_seconds)
            records.extend(parse_sse_directory(body, stock_type=stock_type))
        body = _download(SZSE_URL, {"SHOWTYPE": "xlsx", "CATALOGID": "1110", "TABKEY": "tab1"},
                         referer="https://www.szse.cn/market/product/stock/list/index.html",
                         timeout=self.timeout_seconds)
        records.extend(parse_szse_directory(body))
        return SecurityDirectory(tuple(sorted(_checked_records(records), key=lambda item: item.instrument_id)),
                                 DIRECTORY_SOURCE, datetime.now(timezone.utc), is_fixture=False)
