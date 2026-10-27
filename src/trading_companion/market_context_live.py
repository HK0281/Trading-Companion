"""Opt-in public exchange metadata acquisition with explicit observation semantics."""

import argparse
from dataclasses import asdict
from datetime import date, datetime, timezone
from io import BytesIO
from html.parser import HTMLParser
import json
import os
from pathlib import Path
import re
import tempfile
from urllib.parse import urlencode

from trading_companion.errors import DataSourceError, DataValidationError, InvalidRequestError
from trading_companion.exchange_directory import (
    DIRECTORY_SOURCE, SSE_QUERY, SSE_URL, SZSE_URL, _download, _timeout,
    parse_sse_directory, parse_szse_directory,
)
from trading_companion.market_context import (
    LIVE_SNAPSHOT_VERSION, MAX_SNAPSHOT_BYTES, MarketContextSnapshot, MarketSourceRecord,
    SecurityState, SHANGHAI, load_market_context, parse_market_context_bytes,
)
from trading_companion.security_directory import Security, SecurityDirectory


SOURCE_VERSION = "exchange-market-context-v1"
RISK_URL = "https://query.sse.com.cn/commonSoaQuery.do"
SZSE_JSON_URL = "https://www.szse.cn/api/report/ShowReport/data"
RISK_QUERIES = (("PL_SSGSXX_FXJSBGPLB", {"domesticIndicator": "S", "productType": "0"}, "sse_main_risk"),
                ("SSE_PL_SSGSXX_KCBFXTS_L", {"type": "0"}, "sse_star_risk"))
SSE_PAGE = "https://www.sse.com.cn/assortment/stock/list/share/"
SZSE_PAGE = "https://www.szse.cn/market/product/stock/list/index.html"
RISK_PAGE = "https://www.sse.com.cn/disclosure/listedinfo/riskplate/"


def _listing_date(value):
    if value in (None, "", "-"):
        return None
    if type(value) is datetime:
        if value.time().isoformat() != "00:00:00":
            raise DataValidationError("Unexpected exchange listing timestamp")
        return value.date()
    if type(value) is date:
        return value
    if type(value) is not str:
        raise DataValidationError("Invalid exchange listing date")
    if re.fullmatch(r"[0-9]{8}", value):
        value = value[:4] + "-" + value[4:6] + "-" + value[6:]
    try:
        if re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", value) is None:
            raise ValueError
        return date.fromisoformat(value)
    except ValueError:
        raise DataValidationError("Invalid exchange listing date") from None


def parse_sse_metadata(body: bytes, *, stock_type: str):
    """Reuse complete-page identity validation before reading typed extra columns."""
    securities = parse_sse_directory(body, stock_type=stock_type)
    rows = json.loads(body.decode("utf-8"))["result"]
    board = "star" if stock_type == "8" else "sse_main"
    return tuple((security, board, _listing_date(row.get("LIST_DATE")))
                 for security, row in zip(securities, rows))


def parse_szse_metadata(body: bytes):
    """Board comes from the export's board column, never a code-prefix heuristic."""
    securities = parse_szse_directory(body)
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(body), read_only=True, data_only=False, keep_links=False)
    try:
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_col=100, values_only=True)
        header = next(iterator)
        for field in ("板块", "A股上市日期"):
            if header.count(field) != 1:
                raise DataValidationError("SZSE metadata is missing unique board/listing columns")
        board_index, listing_index = header.index("板块"), header.index("A股上市日期")
        rows = list(iterator)
        if len(rows) != len(securities):
            raise DataValidationError("SZSE metadata row count changed")
        return tuple((security, {"主板": "szse_main", "创业板": "chinext"}.get(row[board_index], "unknown"),
                      _listing_date(row[listing_index])) for security, row in zip(securities, rows))
    finally:
        workbook.close()


class _OfficialName(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag not in {"a", "u"}:
            raise DataValidationError("Unexpected markup in the official security name")

    def handle_endtag(self, tag):
        if tag not in {"a", "u"}:
            raise DataValidationError("Unexpected markup in the official security name")

    def handle_data(self, data):
        self.parts.append(data)


def parse_szse_security_metadata(body: bytes, *, instrument_id: str, observation_date: date):
    """Validate a complete one-stock official report, including its declared report date."""
    try:
        payload = json.loads(body.decode("utf-8"))
        if type(payload) is not list or not payload:
            raise ValueError
        fragments = [item for item in payload if type(item) is dict and type(item.get("metadata")) is dict
                     and item["metadata"].get("tabkey") == "tab1"]
        if len(fragments) != 1:
            raise ValueError
        metadata, rows = fragments[0]["metadata"], fragments[0]["data"]
        reported = _listing_date(str(metadata.get("subname", "")).strip())
        if (metadata.get("catalogid") != "1110" or metadata.get("pageno") != 1
                or metadata.get("pagecount") != 1 or metadata.get("recordcount") != 1
                or any(type(metadata.get(field)) is not int for field in ("pageno", "pagecount", "recordcount"))
                or type(rows) is not list or len(rows) != 1
                or reported is None or not 0 <= (observation_date - reported).days <= 7):
            raise ValueError
        row = rows[0]
        if type(row) is not dict or row.get("agdm") != instrument_id[:6] or not instrument_id.endswith(".SZ"):
            raise ValueError
        if type(row.get("agjc")) is not str or len(row["agjc"]) > 1000:
            raise ValueError
        parser = _OfficialName()
        parser.feed(row["agjc"])
        parser.close()
        security = Security(instrument_id, "".join(parser.parts))
        return ((security, {"主板": "szse_main", "创业板": "chinext"}.get(row.get("bk"), "unknown"),
                 _listing_date(row.get("agssrq"))),)
    except (KeyError, TypeError, ValueError, UnicodeError, RecursionError):
        raise DataValidationError("Invalid, stale or ambiguous SZSE single-security metadata") from None


def parse_sse_risk_status(body: bytes, *, query: str) -> dict[str, str]:
    """Only positive membership in an official risk list establishes a current risk marker."""
    try:
        payload = json.loads(body.decode("utf-8"))
        rows, page = payload["result"], payload["pageHelp"]
        if (payload.get("sqlId") != query or payload.get("actionErrors") != []
                or payload.get("fieldErrors") != {} or type(rows) is not list
                or type(page) is not dict or type(page.get("total")) is not int
                or not 0 <= page["total"] <= 10000 or len(rows) != page["total"]
                or type(page.get("pageNo")) is not int or type(page.get("pageCount")) is not int
                or page.get("pageNo") != 1 or page.get("pageCount") != (1 if rows else 0)):
            raise ValueError
        result = {}
        star = query == "SSE_PL_SSGSXX_KCBFXTS_L"
        if query not in {item[0] for item in RISK_QUERIES}:
            raise ValueError
        for row in rows:
            code = row.get("secCode" if star else "INSTRUMENT_ID")
            name = row.get("secNameCn" if star else "INSTRUMENT_SHORT")
            if type(code) is not str or re.fullmatch(r"[0-9]{6}", code) is None or type(name) is not str:
                raise ValueError
            if not star and code.startswith("900"):
                continue
            if not code.startswith(("688", "689") if star else ("600", "601", "603", "605")):
                raise ValueError
            identifier = code + ".SH"
            if identifier in result:
                raise ValueError
            # Membership proves a risk flag; the official short-name marker refines its category.
            result[identifier] = "star_st" if name.upper().startswith("*ST") else "st" if name.upper().startswith("ST") else "unknown"
        return result
    except (KeyError, TypeError, ValueError, UnicodeError, RecursionError):
        raise DataValidationError("Invalid or incomplete SSE risk-status response") from None


class ExchangeMarketContextSource:
    """Fixed endpoints, injected transport, no credentials, retries or provider substitutions."""

    def __init__(self, *, timeout_seconds=15.0, transport=None, clock=None):
        _timeout(timeout_seconds)
        self.timeout_seconds = timeout_seconds
        self.transport = transport or _download
        self.clock = clock or (lambda: datetime.now(timezone.utc))
        self.last_unavailable = ()

    def read_snapshot(self, *, directory: SecurityDirectory | None = None,
                      instrument_ids: tuple[str, ...] | None = None) -> MarketContextSnapshot:
        if instrument_ids is not None:
            if (type(instrument_ids) is not tuple or not 1 <= len(instrument_ids) <= 3
                    or any(type(identifier) is not str for identifier in instrument_ids)
                    or len(set(instrument_ids)) != len(instrument_ids)):
                raise InvalidRequestError("Select one to three distinct directory-grounded instruments")
            for identifier in instrument_ids:
                Security(identifier, "lookup")
                if not identifier.endswith((".SH", ".SZ")):
                    raise InvalidRequestError("Online market context currently covers SSE and SZSE only")
        requested = set(instrument_ids) if instrument_ids is not None else None
        started = self.clock()
        if type(started) is not datetime or started.tzinfo is None or started.utcoffset() is None:
            raise InvalidRequestError("Market context acquisition requires an aware clock")
        report_date = started.astimezone(SHANGHAI).date()
        fragments, reported_dates = [], {}
        for stock_type, reference in (("1", "sse_main_listing"), ("8", "sse_star_listing")):
            if requested is not None and not any(identifier.endswith(".SH") for identifier in requested):
                continue
            parameters = {"STOCK_TYPE": stock_type, "REG_PROVINCE": "", "CSRC_CODE": "", "STOCK_CODE": "",
                "sqlId": SSE_QUERY, "COMPANY_STATUS": "2,4,5,7,8", "type": "inParams", "isPagination": "true",
                "pageHelp.cacheSize": "1", "pageHelp.beginPage": "1", "pageHelp.pageSize": "10000",
                "pageHelp.pageNo": "1", "pageHelp.endPage": "1"}
            body = self.transport(SSE_URL, parameters, referer=SSE_PAGE, timeout=self.timeout_seconds)
            rows = parse_sse_metadata(body, stock_type=stock_type)
            fragments.append((reference, "sse", SSE_PAGE,
                              tuple(row for row in rows if requested is None or row[0].instrument_id in requested)))
        if requested is None:
            body = self.transport(SZSE_URL, {"SHOWTYPE": "xlsx", "CATALOGID": "1110", "TABKEY": "tab1"},
                                  referer=SZSE_PAGE, timeout=self.timeout_seconds)
            fragments.append(("szse_listing", "szse", SZSE_PAGE, parse_szse_metadata(body)))
        else:
            for identifier in sorted(requested):
                if not identifier.endswith(".SZ"):
                    continue
                parameters = {"SHOWTYPE": "JSON", "CATALOGID": "1110", "TABKEY": "tab1", "txtDMorJC": identifier[:6], "PAGENO": "1"}
                body = self.transport(SZSE_JSON_URL, parameters,
                    referer=SZSE_PAGE, timeout=self.timeout_seconds)
                reference = "szse_listing_" + identifier[:6]
                rows = parse_szse_security_metadata(body, instrument_id=identifier, observation_date=report_date)
                reported_dates[reference] = _listing_date(next(item["metadata"]["subname"].strip()
                    for item in json.loads(body) if item.get("metadata", {}).get("tabkey") == "tab1"))
                fragments.append((reference, "szse", SZSE_JSON_URL + "?" + urlencode(parameters), rows))
        risk, missing = {}, ["exchange_calendar_not_fetched", "suspension_status_not_fetched", "szse_risk_status_not_fetched"]
        risk_references = []
        for query, parameters, reference in RISK_QUERIES:
            if requested is not None and not any(identifier.endswith(".SH") for identifier in requested):
                continue
            try:
                body = self.transport(RISK_URL, {"sqlId": query, **parameters}, referer=RISK_PAGE,
                                      timeout=self.timeout_seconds)
                statuses = parse_sse_risk_status(body, query=query)
                risk.update({identifier: (status, reference) for identifier, status in statuses.items()})
                risk_references.append(reference)
            except (DataSourceError, DataValidationError):
                missing.append(reference + "_unavailable")
        now = self.clock()
        if type(now) is not datetime or now.tzinfo is None or now.utcoffset() is None:
            raise InvalidRequestError("Market context acquisition requires an aware clock")
        today = now.astimezone(SHANGHAI).date()
        if today != report_date:
            raise DataValidationError("Market-context acquisition crossed its observation date boundary")
        observed = {security.instrument_id: security.name for *_, rows in fragments for security, _, _ in rows}
        if requested is not None and set(observed) != requested:
            raise DataValidationError("Requested market-context identities are not completely covered")
        if directory is not None:
            if type(directory) is not SecurityDirectory or directory.source != DIRECTORY_SOURCE:
                raise DataValidationError("Market context requires an exchange-validated directory")
            directory.assert_usable(today, development=False)
            expected = {security.instrument_id: security.name for security in directory.securities
                        if requested is None or security.instrument_id in requested}
            if expected != observed:
                raise DataValidationError("Directory and metadata identities differ; refresh both snapshots")
        states, sources = [], []
        for reference, source, page, rows in fragments:
            source_date = reported_dates.get(reference, today)
            sources.append(MarketSourceRecord(reference, source, page, source_date, now, None, "unknown",
                ("board", "listing_date"), "selected_current_a_shares" if requested else "current_listed_a_shares"))
            for security, board, listing_date in rows:
                status, st_reference = risk.get(security.instrument_id, ("unknown", None))
                states.append(SecurityState(security.instrument_id, source_date, source_date, board, listing_date,
                    status, "unknown", None, source, reference, now, st_reference))
        for reference in risk_references:
            query, parameters, _ = next(item for item in RISK_QUERIES if item[2] == reference)
            source_url = RISK_URL + "?" + urlencode({"sqlId": query, **parameters})
            sources.append(MarketSourceRecord(reference, "sse", source_url, today, now, None, "unknown",
                ("st_status",), "positive_risk_membership_only"))
        self.last_unavailable = tuple(missing)
        return MarketContextSnapshot("live-" + now.strftime("%Y%m%dT%H%M%S%f"), now, False,
                                     tuple(sorted(states, key=lambda item: item.instrument_id)), (), tuple(sources))


def _encode(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    raise TypeError("Unsupported market-context value")


def save_market_context(snapshot: MarketContextSnapshot, path: str | Path) -> None:
    """Atomically publish a validated live snapshot without replacing fixtures or unrelated JSON."""
    selected, temporary = Path(path), None
    if selected.suffix.lower() != ".json" or selected.is_symlink() or snapshot.is_fixture:
        raise InvalidRequestError("Live context output must be a non-symlink JSON snapshot")
    if selected.exists() and load_market_context(selected).is_fixture:
        raise InvalidRequestError("Refusing to overwrite a fixture market context")
    encoded = (json.dumps({"version": LIVE_SNAPSHOT_VERSION, **asdict(snapshot)}, default=_encode,
                          ensure_ascii=False, allow_nan=False, indent=2) + "\n").encode("utf-8")
    if len(encoded) > MAX_SNAPSHOT_BYTES or parse_market_context_bytes(encoded) != snapshot:
        raise DataValidationError("Market context failed a bounded serialization round trip")
    try:
        selected.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(mode="wb", dir=selected.parent, prefix=".market-context-",
                                         suffix=".json", delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(encoded)
            handle.flush()
            os.fsync(handle.fileno())
        if load_market_context(temporary) != snapshot:
            raise DataValidationError("Saved market context did not round-trip")
        os.replace(temporary, selected)
        temporary = None
    except OSError:
        raise DataSourceError("Cannot save the market-context snapshot") from None
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def refresh_market_context(path, *, source=None, directory=None, instrument_ids=None):
    adapter = source or ExchangeMarketContextSource()
    snapshot = adapter.read_snapshot(directory=directory, instrument_ids=instrument_ids)
    save_market_context(snapshot, path)
    return snapshot


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Refresh current SSE/SZSE metadata from fixed public exchange endpoints")
    parser.add_argument("--live", action="store_true", help="Enable three listing and two optional SSE risk HTTP reads")
    parser.add_argument("--output", default="data/market_context_local.json")
    parser.add_argument("--timeout", type=float, default=15.0)
    parser.add_argument("--instrument", action="append", help="Fetch one to three explicit stocks; SZSE uses a filtered official JSON report")
    arguments = parser.parse_args(argv)
    if not arguments.live:
        print(json.dumps({"status": "not_requested", "network_requested": False, "snapshot_replaced": False}))
        return 0
    try:
        source = ExchangeMarketContextSource(timeout_seconds=arguments.timeout)
        snapshot = refresh_market_context(arguments.output, source=source,
                                          instrument_ids=tuple(arguments.instrument) if arguments.instrument else None)
    except (DataSourceError, DataValidationError, InvalidRequestError, OSError):
        print(json.dumps({"status": "failed", "code": "market_context_refresh_failed", "snapshot_replaced": False}))
        return 1
    print(json.dumps({"status": "limited", "code": "current_market_context_refreshed", "snapshot_replaced": True,
        "source_version": SOURCE_VERSION, "snapshot_id": snapshot.snapshot_id,
        "retrieved_at": snapshot.retrieved_at.isoformat(), "securities": len(snapshot.security_states),
        "source_records": len(snapshot.source_records), "unavailable": list(source.last_unavailable),
        "historical_status_available": False, "network_requested": True, "model_calls": 0}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
